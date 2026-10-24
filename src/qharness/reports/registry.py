from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import yaml
from docxtpl import DocxTemplate
from jsonschema import Draft202012Validator
from openpyxl import load_workbook

from qharness.reports.ooxml import inspect_ooxml
from qharness.reports.xlsx import named_cell


class InvalidReportContract(ValueError):
    pass


class ChecksumDrift(RuntimeError):
    pass


@dataclass(frozen=True, slots=True)
class ReportContract:
    path: Path
    front_matter: dict[str, Any]
    body: str


@dataclass(frozen=True, slots=True)
class RegisteredTemplate:
    contract: ReportContract
    template_path: Path
    contract_sha256: str
    template_sha256: str
    license: str
    trusted: bool


def _sha256(path: Path) -> str:
    if any(part.is_symlink() for part in (path, *path.parents)):
        raise InvalidReportContract("report paths must not contain symlinks")
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_contract(path: Path) -> ReportContract:
    content = path.read_text(encoding="utf-8")
    match = re.match(r"\A---\s*\n(.*?)\n---\s*\n(.*)\Z", content, re.S)
    if not match:
        raise InvalidReportContract("report.md must contain YAML front matter")
    front_matter = yaml.safe_load(match.group(1))
    if not isinstance(front_matter, dict):
        raise InvalidReportContract("report.md front matter must be an object")
    schema_path = Path(__file__).with_name("report-contract.schema.json")
    schema = yaml.safe_load(schema_path.read_text(encoding="utf-8"))
    errors = sorted(
        Draft202012Validator(schema).iter_errors(front_matter), key=lambda item: list(item.path)
    )
    if errors:
        raise InvalidReportContract(errors[0].message)
    required_headings = ("# Purpose", "# Data Rules", "# Rendering Rules", "# Validation Rules")
    if any(heading not in match.group(2) for heading in required_headings):
        raise InvalidReportContract("report.md is missing a required heading")
    return ReportContract(path.resolve(), front_matter, match.group(2))


class ReportRegistry:
    def register(
        self, contract_path: Path, *, trusted: bool, license_name: str
    ) -> RegisteredTemplate:
        if not trusted:
            raise InvalidReportContract("templates must be operator-trusted")
        if not license_name.strip():
            raise InvalidReportContract("template license is required")
        contract_hash = _sha256(contract_path.absolute())
        contract = load_contract(contract_path)
        template_path = contract.path.parent / contract.front_matter["template_file"]
        try:
            template_path.relative_to(contract.path.parent)
        except ValueError as exc:
            raise InvalidReportContract("template path escapes the contract root") from exc
        expected_suffix = f".{contract.front_matter['output_format']}"
        if template_path.suffix.lower() != expected_suffix or not template_path.is_file():
            raise InvalidReportContract("declared template is missing or has the wrong format")
        if contract.front_matter["required_permission"] != "report:generate":
            raise InvalidReportContract("unsupported report permission")
        template_hash = _sha256(template_path)
        try:
            inspect_ooxml(template_path)
            declared = {
                field
                for section in contract.front_matter["required_sections"]
                for field in section.get("required_fields", [])
            }
            if expected_suffix == ".docx":
                variables = DocxTemplate(template_path).get_undeclared_template_variables()
                if variables - declared:
                    raise InvalidReportContract(
                        f"undeclared template variables: {sorted(variables - declared)}"
                    )
            else:
                workbook = load_workbook(template_path, data_only=False)
                validation = contract.front_matter["validation"]
                if any(
                    sheet not in workbook.sheetnames
                    for sheet in validation.get("required_sheets", [])
                ):
                    raise InvalidReportContract("required worksheet is missing")
                if any(
                    name not in workbook.defined_names
                    for name in validation.get("required_named_ranges", [])
                ):
                    raise InvalidReportContract("required named range is missing")
                for name in declared:
                    if name not in workbook.defined_names:
                        raise InvalidReportContract(
                            f"required field has no input named range: {name}"
                        )
                for name in declared | set(validation.get("required_named_ranges", [])):
                    named_cell(workbook, name)
                workbook.close()
        except InvalidReportContract:
            raise
        except Exception as exc:
            raise InvalidReportContract("unsupported or invalid OOXML template") from exc
        if _sha256(contract.path) != contract_hash or _sha256(template_path) != template_hash:
            raise ChecksumDrift("report contract or template changed during preflight")
        return RegisteredTemplate(
            contract,
            template_path,
            contract_hash,
            template_hash,
            license_name.strip(),
            trusted,
        )

    def verify(self, record: RegisteredTemplate) -> None:
        if (
            _sha256(record.contract.path) != record.contract_sha256
            or _sha256(record.template_path) != record.template_sha256
        ):
            raise ChecksumDrift("report contract or template checksum changed")
