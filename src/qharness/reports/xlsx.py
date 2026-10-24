from __future__ import annotations

import re
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.workbook.workbook import Workbook

from qharness.reports.ooxml import inspect_ooxml


class ReportValidationError(ValueError):
    pass


def _cell(workbook: Workbook, reference: str) -> Cell:
    sheet_name, cell = reference.split("!", 1)
    if sheet_name not in workbook.sheetnames:
        raise ReportValidationError(f"missing sheet: {sheet_name}")
    result = workbook[sheet_name][cell]
    if not isinstance(result, Cell):
        raise ReportValidationError("input must reference a single non-merged cell")
    return result


def named_cell(workbook: Workbook, name: str) -> str:
    definition = workbook.defined_names[name]
    destinations = list(definition.destinations)
    if definition.type != "RANGE" or len(destinations) != 1:
        raise ReportValidationError("named input must reference one static cell")
    sheet, coordinate = destinations[0]
    if ":" in coordinate:
        raise ReportValidationError("named input must reference one static cell")
    reference = f"{sheet}!{coordinate}"
    _cell(workbook, reference)
    return reference


def _sum_formula(workbook: Workbook, sheet_name: str, formula: str) -> float | None:
    match = re.fullmatch(r"=SUM\(([A-Z]+\d+):([A-Z]+\d+)\)", formula, re.I)
    if not match:
        return None
    values = [
        cell.value
        for row in workbook[sheet_name][f"{match.group(1)}:{match.group(2)}"]
        for cell in row
        if isinstance(cell.value, (int, float))
    ]
    return float(sum(values))


def render_xlsx(
    template: Path,
    output: Path,
    values: dict[str, object],
    required_sheets: list[str],
    expected_totals: dict[str, float],
) -> None:
    try:
        inspect_ooxml(template)
        workbook = load_workbook(template, data_only=False)
        if any(sheet not in workbook.sheetnames for sheet in required_sheets):
            raise ReportValidationError("required worksheet is missing")
        original_formulas = {
            f"{sheet.title}!{cell.coordinate}": cell.value
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        }
        for reference, value in values.items():
            cell = _cell(workbook, reference)
            if isinstance(cell.value, str) and cell.value.startswith("="):
                raise ReportValidationError("template formulas cannot be overwritten")
            if value is not None and not isinstance(
                value, (str, int, float, Decimal, date, datetime, time, timedelta)
            ):
                raise ReportValidationError("workbook input must be a scalar")
            cell.value = value
            if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
                cell.data_type = "s"
        for reference, expected in expected_totals.items():
            cell = _cell(workbook, reference)
            calculated = _sum_formula(workbook, reference.split("!", 1)[0], str(cell.value))
            if calculated is None or calculated != float(expected):
                raise ReportValidationError("independent total validation failed")
        workbook.save(output)
        reopened = load_workbook(output, data_only=False)
        for reference, formula in original_formulas.items():
            if _cell(reopened, reference).value != formula:
                raise ReportValidationError("formula changed during rendering")
        if any(
            isinstance(cell.value, str) and ("{{" in cell.value or "{%" in cell.value)
            for sheet in reopened
            for row in sheet
            for cell in row
        ):
            raise ReportValidationError("unresolved workbook placeholder")
        reopened.close()
        workbook.close()
    except ReportValidationError:
        raise
    except Exception as exc:
        raise ReportValidationError("XLSX rendering or reopen failed") from exc


def render_named_workbook(
    template: Path, output: Path, context: dict[str, object], validation: dict[str, Any]
) -> None:
    workbook = load_workbook(template, data_only=False)
    values: dict[str, object] = {}
    for name, value in context.items():
        if name not in workbook.defined_names:
            continue
        values[named_cell(workbook, name)] = value
    totals: dict[str, float] = {}
    test_total = context.get("test_total")
    if "test_total" in workbook.defined_names and isinstance(test_total, (int, float)):
        reference = named_cell(workbook, "test_total")
        values.pop(reference, None)
        totals[reference] = float(test_total)
    for name in validation.get("required_named_ranges", []):
        if name not in workbook.defined_names:
            raise ReportValidationError("required named range is missing")
    workbook.close()
    render_xlsx(template, output, values, validation.get("required_sheets", []), totals)
