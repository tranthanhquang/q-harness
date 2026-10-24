from pathlib import Path

import pytest

from qharness.reports.registry import ChecksumDrift, ReportRegistry


@pytest.mark.parametrize("license_name,trusted", [(" ", True), ("MIT", False)])
def test_registry_requires_explicit_trust_and_license(tmp_path, license_name, trusted):
    from qharness.reports.registry import InvalidReportContract

    contract, _ = build_minimal_contract(tmp_path)
    with pytest.raises(InvalidReportContract):
        ReportRegistry().register(contract, trusted=trusted, license_name=license_name)


def test_registry_rejects_symlinked_directory(tmp_path):
    from qharness.reports.registry import InvalidReportContract

    real = tmp_path / "real"
    real.mkdir()
    build_minimal_contract(real)
    alias = tmp_path / "alias"
    alias.symlink_to(real, target_is_directory=True)
    with pytest.raises(InvalidReportContract, match="symlink"):
        ReportRegistry().register(alias / "report.md", trusted=True, license_name="MIT")


def test_registry_detects_template_change_during_preflight(tmp_path, monkeypatch):
    import qharness.reports.registry as registry

    contract, template = build_minimal_contract(tmp_path)
    inspect = registry.inspect_ooxml

    def replace_after_inspection(path):
        inspect(path)
        template.write_bytes(template.read_bytes() + b"changed")

    monkeypatch.setattr(registry, "inspect_ooxml", replace_after_inspection)
    with pytest.raises(ChecksumDrift):
        ReportRegistry().register(contract, trusted=True, license_name="MIT")


@pytest.mark.parametrize("reference", ["'Tests'!$A$1:$A$2", "42", "'Missing'!$A$1"])
def test_registry_rejects_non_static_single_input_cells(tmp_path, reference):
    from openpyxl import Workbook
    from openpyxl.workbook.defined_name import DefinedName

    from qharness.reports.registry import InvalidReportContract

    contract, _ = build_minimal_contract(tmp_path)
    contract.write_text(contract.read_text().replace("docx", "xlsx"))
    workbook = Workbook()
    workbook.active.title = "Tests"
    workbook.defined_names.add(DefinedName("project_name", attr_text=reference))
    workbook.save(tmp_path / "template.xlsx")
    with pytest.raises(InvalidReportContract):
        ReportRegistry().register(contract, trusted=True, license_name="MIT")


@pytest.mark.integration
def test_registry_pins_contract_and_template_checksums(tmp_path: Path) -> None:
    contract, template = build_minimal_contract(tmp_path)
    record = ReportRegistry().register(contract, trusted=True, license_name="MIT")
    assert record.template_sha256 and record.contract_sha256
    template.write_bytes(template.read_bytes() + b"drift")
    with pytest.raises(ChecksumDrift):
        ReportRegistry().verify(record)


def test_registry_rejects_undeclared_docx_variables(tmp_path: Path) -> None:
    from docx import Document

    from qharness.reports.registry import InvalidReportContract

    contract, template = build_minimal_contract(tmp_path)
    document = Document(template)
    document.add_paragraph("{{ undeclared_secret }}")
    document.save(template)
    with pytest.raises(InvalidReportContract, match="undeclared"):
        ReportRegistry().register(contract, trusted=True, license_name="MIT")


def build_minimal_contract(root: Path) -> tuple[Path, Path]:
    from docx import Document

    template = root / "template.docx"
    document = Document()
    document.add_heading("{{ project_name }}", level=1)
    document.save(template)
    contract = root / "report.md"
    contract.write_text(
        """---
template_id: weekly-progress
version: 1.0.0
scope: project
output_format: docx
template_file: template.docx
language: vi
required_permission: report:generate
data_sources: [project]
required_sections:
  - id: overview
    title: Overview
    required_fields: [project_name]
missing_data_policy: fail
validation: {openable: true, no_unresolved_placeholders: true, preserve_template_checksum: true}
---
# Purpose
Test.
# Data Rules
Project only.
# Rendering Rules
Strict.
# Validation Rules
Openable.
""",
        encoding="utf-8",
    )
    return contract, template
