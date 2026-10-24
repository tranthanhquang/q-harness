from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.workbook.defined_name import DefinedName

from qharness.reports.xlsx import ReportValidationError, render_named_workbook, render_xlsx


@pytest.mark.integration
def test_xlsx_preserves_formula_and_validates_total(tmp_path: Path) -> None:
    template = tmp_path / "template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tests"
    sheet["A2"] = 2
    sheet["A3"] = 3
    sheet["A4"] = "=SUM(A2:A3)"
    workbook.save(template)
    output = tmp_path / "output.xlsx"
    render_xlsx(template, output, {"Tests!B1": "Run"}, ["Tests"], {"Tests!A4": 5})
    reopened = load_workbook(output, data_only=False)
    assert reopened["Tests"]["A4"].value == "=SUM(A2:A3)"
    with pytest.raises(ReportValidationError):
        render_xlsx(template, output, {}, ["Missing"], {})


def test_named_inputs_preserve_styles_escape_formula_text_and_check_totals(tmp_path: Path):
    template, output = tmp_path / "template.xlsx", tmp_path / "output.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tests"
    sheet["B1"].font = Font(bold=True, color="FF0000")
    sheet["B2"], sheet["B3"], sheet["B4"] = 2, 3, "=SUM(B2:B3)"
    for name, cell in (("project_name", "B1"), ("test_total", "B4")):
        workbook.defined_names.add(DefinedName(name, attr_text=f"'Tests'!${cell[0]}${cell[1:]}"))
    workbook.save(template)
    validation = {"required_sheets": ["Tests"], "required_named_ranges": ["project_name"]}
    render_named_workbook(template, output, {"project_name": "=1+1", "test_total": 5}, validation)
    reopened = load_workbook(output, data_only=False)
    assert reopened["Tests"]["B1"].data_type == "s"
    assert reopened["Tests"]["B1"].value == "=1+1"
    assert reopened["Tests"]["B1"].font.bold
    assert reopened["Tests"]["B1"].font.color.rgb == "00FF0000"
    assert reopened["Tests"]["B4"].value == "=SUM(B2:B3)"
    reopened.close()
    with pytest.raises(ReportValidationError, match="independent total"):
        render_named_workbook(template, output, {"test_total": 6}, validation)
    with pytest.raises(ReportValidationError, match="formulas cannot"):
        render_xlsx(template, output, {"Tests!B4": 5}, ["Tests"], {})
    with pytest.raises(ReportValidationError, match="named range"):
        render_named_workbook(template, output, {}, {"required_named_ranges": ["absent"]})


def test_external_ooxml_relationships_are_rejected(tmp_path: Path):
    template = tmp_path / "external.xlsx"
    workbook = Workbook()
    workbook.active["A1"].hyperlink = "https://example.com"
    workbook.save(template)
    with pytest.raises(ReportValidationError):
        render_xlsx(template, tmp_path / "output.xlsx", {}, [], {})
