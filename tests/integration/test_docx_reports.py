from pathlib import Path

import pytest
from docx import Document

from qharness.reports.docx import ReportValidationError, render_docx


@pytest.mark.integration
def test_docx_render_reopens_and_rejects_unresolved_values(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    document = Document()
    document.add_heading("{{ project_name }}", level=1)
    document.save(template)
    output = tmp_path / "output.docx"
    render_docx(template, output, {"project_name": "Q-Harness"}, ["Q-Harness"])
    assert Document(output).paragraphs[0].text == "Q-Harness"
    with pytest.raises(ReportValidationError):
        render_docx(template, output, {}, ["Q-Harness"])


def test_docx_validates_header_and_footer_placeholders(tmp_path: Path) -> None:
    template, output = tmp_path / "template.docx", tmp_path / "output.docx"
    document = Document()
    document.add_paragraph("Report")
    document.sections[0].header.paragraphs[0].text = "{{ value }}"
    document.sections[0].footer.paragraphs[0].text = "Footer"
    document.save(template)
    with pytest.raises(ReportValidationError):
        render_docx(template, output, {"value": "{{ unresolved }}"}, ["Footer"])
    render_docx(template, output, {"value": "Header"}, ["Header", "Footer"])
