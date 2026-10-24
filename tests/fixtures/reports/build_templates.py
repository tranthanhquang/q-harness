from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName


def build(root: Path) -> None:
    weekly = root / "weekly-progress" / "template.docx"
    weekly.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading("Báo cáo tiến độ", level=1)
    document.add_paragraph("Dự án: {{ project_name }}")
    document.add_paragraph("Số task: {{ task_count }}")
    document.save(weekly)

    tests = root / "test-results" / "template.xlsx"
    tests.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tests"
    sheet["A1"] = "Project"
    for row, (field, label) in enumerate(
        (
            ("project_name", "Project"),
            ("test_passed", "Passed"),
            ("test_failed", "Failed"),
            ("test_not_run", "Not run"),
        ),
        start=1,
    ):
        sheet[f"A{row}"] = label
        workbook.defined_names.add(DefinedName(field, attr_text=f"'Tests'!$B${row}"))
    sheet["A5"] = "Total"
    sheet["B5"] = "=SUM(B2:B4)"
    workbook.defined_names.add(DefinedName("test_total", attr_text="'Tests'!$B$5"))
    workbook.save(tests)


if __name__ == "__main__":
    build(Path(__file__).resolve().parents[3] / "src/qharness/reports/templates")
